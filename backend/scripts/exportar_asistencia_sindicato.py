"""Consolidado de la jornada de UNA entidad: mismo formato del Excel de inscripción + columna de asistencia.

Uso (desde backend/, con el .venv): python scripts/exportar_asistencia_sindicato.py "<ENTIDAD>[|<VARIANTE>]" salida.xlsx
Cruza personal (por entidad) con pre_test (inscripción) y asistencia (Pi) de la campaña 2026.
Sin inscripción -> "No inscrito aún"; sin asistencia -> "Sin asistencia aún". Hora en Colombia (UTC-5).
"""
import asyncio, asyncpg, sys, os
from pathlib import Path
os.chdir(Path(__file__).resolve().parents[1])
from dotenv import dotenv_values
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ENTIDAD, SALIDA = sys.argv[1], sys.argv[2]
NR = "No respondido aún"
NI = "No inscrito aún"
SA = "Sin asistencia aún"
from datetime import date, datetime
from datetime import timezone, timedelta
_BOG = timezone(timedelta(hours=-5))  # Colombia, sin horario de verano
MESES = ["enero","febrero","marzo","abril","mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"]
def fdia(x):
    try:
        d = date.fromisoformat(str(x)[:10]); return f"{d.day} de {MESES[d.month-1]}"
    except Exception: return x
def fhora(x):
    try:
        h = datetime.strptime(str(x).strip()[:5].rstrip(":"), "%H:%M"); return f"{h.hour%12 or 12}:{h.minute:02d} {'AM' if h.hour<12 else 'PM'}"
    except Exception: return x
def fasis(t):
    t = t.astimezone(_BOG)
    return (f"{t.day} de {MESES[t.month-1]}", f"{t.hour%12 or 12}:{t.minute:02d} {'AM' if t.hour<12 else 'PM'}")
SQL = """
select p.cedula, p.nombre_completo, p.cargo, p.proceso, p.telefono,
       coalesce(p.email_institucional, p.email_secundario) as email,
       t.puntuacion, t.dia, t.horario, t.salon, a.creado_en as asistio_en
from personal p
left join lateral (
  select * from pre_test t where t.cedula = p.cedula and t.campana = '2026'
  order by t.creado_en desc limit 1) t on true
left join lateral (
  select creado_en from asistencia a where a.cedula = p.cedula::text and a.campana = '2026'
  order by a.creado_en limit 1) a on true
where upper(p.entidad) = any($1::text[])
order by p.nombre_completo
"""
async def main():
    c = await asyncpg.connect(dotenv_values(".env")["DATABASE_URL"], statement_cache_size=0)
    rows = await c.fetch(SQL, [e.upper() for e in ENTIDAD.split("|")])
    wb = Workbook(); ws = wb.active; ws.title = "Asistencia"
    head = ["Cédula","Nombre","Cargo","Proceso","Teléfono","Email","Día inscrito","Hora inscrita","Salón escogido","Fecha de asistencia","Hora de asistencia"]
    ws.append(head)
    for r in rows:
        v = lambda x: x if x not in (None, "") else NR
        insc = bool(r["dia"])
        ws.append([r["cedula"], r["nombre_completo"], v(r["cargo"]), v(r["proceso"]), v(r["telefono"]), v(r["email"]),
                   v(fdia(r["dia"])) if insc else NI,
                   v(fhora(r["horario"])) if insc else NI, v(r["salon"]) if insc else NI,
                   *(fasis(r["asistio_en"]) if r["asistio_en"] else (SA, SA))])
    for c_ in ws[1]:
        c_.font = Font(bold=True, color="FFFFFF"); c_.fill = PatternFill("solid", fgColor="327531")
        c_.alignment = Alignment(horizontal="center", vertical="center")
    for i, h in enumerate(head, 1):
        w = max(len(str(c_.value or "")) for c_ in ws[get_column_letter(i)]) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(max(w*1.2, len(h)*1.3+4), 55)
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "0"
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    n = len(rows)
    t = n + 3   # fila del encabezado del resumen, fuera del rango del filtro
    ws.cell(t, 2, "Asistencia de la jornada"); ws.cell(t, 3, "Personas")
    ws.cell(t+1, 2, "Con asistencia registrada")
    ws.cell(t+1, 3, f"=COUNTIF(J2:J{n+1},\"<>{SA}\")")
    ws.cell(t+2, 2, "Sin asistencia aún")
    ws.cell(t+2, 3, f"=COUNTIF(J2:J{n+1},\"{SA}\")")
    ws.cell(t+3, 2, "Total"); ws.cell(t+3, 3, f"=C{t+1}+C{t+2}")
    for col in (2, 3):
        h = ws.cell(t, col); h.font = Font(bold=True, color="FFFFFF"); h.fill = PatternFill("solid", fgColor="327531")
        ws.cell(t+3, col).font = Font(bold=True)
    ws.cell(t, 3).alignment = Alignment(horizontal="center")
    for r_ in range(t+1, t+4): ws.cell(r_, 3).alignment = Alignment(horizontal="center")
    from openpyxl.chart import PieChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
    ch = PieChart(); ch.title = "Asistencia: registrada vs. pendiente"
    ch.add_data(Reference(ws, min_col=3, min_row=t, max_row=t+2), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=t+1, max_row=t+2))
    for i, col in enumerate(["76B82A", "D5DAE6"]):
        pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = col
        pt.graphicalProperties.line.solidFill = "FFFFFF"; pt.graphicalProperties.line.width = 28575
        ch.series[0].dPt.append(pt)
    dl = DataLabelList(); dl.showVal = True; dl.showPercent = True
    dl.showCatName = False; dl.showSerName = False; dl.showLeaderLines = False; dl.separator = " · "
    cp = CharacterProperties(sz=1200, b=True, solidFill="2B2B2B")
    dl.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
    ch.dataLabels = dl
    ch.legend.position = "b"
    ch.height = 9; ch.width = 15
    ws.add_chart(ch, f"B{t+5}")
    wb.save(SALIDA)
    print(len(rows), "filas")
    for r in rows: print(r["cedula"], r["dia"], r["horario"], r["asistio_en"])
asyncio.run(main())
