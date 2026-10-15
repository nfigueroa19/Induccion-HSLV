"""Excel offline con TODAS las personas inscritas (pre_test 2026): nombre, cédula, cargo, proceso, entidad, inscripción.

Uso (desde backend/): python scripts/exportar_inscritos.py salida.xlsx [limite]
Trae datos personales: guardar fuera del repo y del vault.
"""
import asyncio, asyncpg, sys, os
from pathlib import Path
os.chdir(Path(__file__).resolve().parents[1])
from dotenv import dotenv_values
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import date, datetime
SALIDA = sys.argv[1]; LIMITE = int(sys.argv[2]) if len(sys.argv) > 2 else None
NR = "Sin dato"
MESES = ["enero","febrero","marzo","abril","mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"]
def fdia(x):
    d = date.fromisoformat(str(x)[:10]); return f"{d.day} de {MESES[d.month-1]}"
def fhora(x):
    h = datetime.strptime(str(x).strip()[:5].rstrip(":"), "%H:%M"); return f"{h.hour%12 or 12}:{h.minute:02d} {'AM' if h.hour<12 else 'PM'}"
SQL = """
select p.cedula, p.nombre_completo, p.cargo, p.proceso, p.entidad, t.dia, t.horario, t.salon
from pre_test t
left join personal p on p.cedula = t.cedula
where t.campana = '2026' and t.salon is not null
order by coalesce(p.nombre_completo, 'zzz'), t.cedula
""" + (f" limit {LIMITE}" if LIMITE else "")
async def main():
    c = await asyncpg.connect(dotenv_values(".env")["DATABASE_URL"], statement_cache_size=0)
    rows = await c.fetch(SQL)
    wb = Workbook(); ws = wb.active; ws.title = "Inscritos"
    head = ["Nombre", "Cédula", "Cargo", "Proceso", "Entidad", "Inscripción"]
    ws.append(head)
    v = lambda x: x if x not in (None, "") else NR
    for r in rows:
        insc = " · ".join(str(x) for x in (fdia(r["dia"]) if r["dia"] else None,
               fhora(r["horario"]) if r["horario"] else None, r["salon"]) if x) or NR
        ws.append([v(r["nombre_completo"]), r["cedula"], v(r["cargo"]), v(r["proceso"]), v(r["entidad"]), insc])
    for c_ in ws[1]:
        c_.font = Font(bold=True, color="FFFFFF"); c_.fill = PatternFill("solid", fgColor="327531")
        c_.alignment = Alignment(horizontal="center", vertical="center")
    for i, h in enumerate(head, 1):
        w = max(len(str(c_.value or "")) for c_ in ws[get_column_letter(i)]) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(max(w*1.2, len(h)*1.3+4), 55)
    for row in ws.iter_rows(min_row=2): row[1].number_format = "0"
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(SALIDA)
    print(len(rows), "filas")
    for r in rows[:5]: print(r["nombre_completo"], r["cedula"], r["dia"], r["horario"], r["salon"])
asyncio.run(main())
