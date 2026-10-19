"""Certificados de la Reinducción 2026 de UNA entidad/sindicato, desde la tabla `certificacion`.

Se certifica a quien tiene asistencia + diagnóstico ADN Susana + post-test.
Genera un .zip con:
  - "Resumen <ENTIDAD>.xlsx": cuántos completaron, cuántos no y por qué, más el detalle por persona.
  - "certificados.zip": un PDF por persona certificada, "<cedula>_<NOMBRE>_reinduccion2026.pdf".

Uso (python del sistema con pymupdf, openpyxl y asyncpg):
    python scripts/generar_certificados_entidad.py "<ENTIDAD>[|<VARIANTE>...]" <carpeta_salida> [plantilla.pdf]
La tabla `certificacion` es una foto: regenerarla antes si cambió la fuente. Hora en Colombia (UTC-5).
Fuentes Poppins (licencia OFL) en scripts/fuentes_certificado/.
"""
import asyncio, io, os, re, sys, unicodedata, zipfile
from datetime import timezone, timedelta
from pathlib import Path

import asyncpg
import fitz
from dotenv import dotenv_values
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import marca_agua_certificado as marca

AQUI = Path(__file__).resolve().parent
BACKEND = AQUI.parent
PLANTILLA = Path(sys.argv[3]) if len(sys.argv) > 3 else BACKEND.parent.parent / "Induccion" / "certificado reinduccion.pdf"
FUENTE_NOMBRE = AQUI / "fuentes_certificado" / "Poppins-ExtraBold.ttf"
FUENTE_CC = AQUI / "fuentes_certificado" / "Poppins-Regular.ttf"
VERDE = (6 / 255, 104 / 255, 57 / 255)   # el mismo verde del texto de la plantilla
BOG = timezone(timedelta(hours=-5))      # Colombia, sin horario de verano
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# Posiciones medidas en la plantilla (792 x 612 pt)
CENTRO_X = 396
NOMBRE_Y, NOMBRE_TAM, NOMBRE_MIN, NOMBRE_ANCHO = 309.9, 20.2, 12.0, 500
CC_X, CC_Y, CC_TAM = CENTRO_X, 340.06, 13   # misma vertical que el centro del nombre (la línea va de 280.8 a 575.0)


def cedula_puntos(c) -> str:
    return f"{int(c):,}".replace(",", ".")


def pdf_certificado(fondo: bytes, nombre: str, cedula) -> bytes:   # `fondo`: JPEG del fondo ya marcado
    d = fitz.open()
    pg = d.new_page(width=792, height=612)
    pg.insert_image(pg.rect, stream=fondo)
    pg.insert_font(fontname="PopXB", fontfile=str(FUENTE_NOMBRE))
    pg.insert_font(fontname="PopR", fontfile=str(FUENTE_CC))
    nombre = " ".join(nombre.upper().split())
    f = fitz.Font(fontfile=str(FUENTE_NOMBRE))
    tam = NOMBRE_TAM
    while tam > NOMBRE_MIN and f.text_length(nombre, fontsize=tam) > NOMBRE_ANCHO:
        tam -= 0.25
    x = CENTRO_X - f.text_length(nombre, fontsize=tam) / 2
    pg.insert_text((x, NOMBRE_Y), nombre, fontname="PopXB", fontsize=tam, color=VERDE)
    cc = cedula_puntos(cedula)
    cc_x = CC_X - fitz.Font(fontfile=str(FUENTE_CC)).text_length(cc, fontsize=CC_TAM) / 2
    pg.insert_text((cc_x, CC_Y), cc, fontname="PopR", fontsize=CC_TAM, color=VERDE)
    d.set_metadata({"title": f"Certificado de asistencia - {nombre}", "author": "Hospital Susana López de Valencia E.S.E."})
    d.subset_fonts()
    return d.tobytes(garbage=4, deflate=True)


def ascii_nombre(s: str) -> str:
    s = unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", "_", s).strip("_")


def fasis(t):
    t = t.astimezone(BOG)
    return f"{t.day} de {MESES[t.month - 1]}", f"{t.hour % 12 or 12}:{t.minute:02d} {'AM' if t.hour < 12 else 'PM'}"


def faltas(r):
    f = []
    if r["asistencia"] is None: f.append("asistencia")
    if r["puntaje_posttest"] is None: f.append("post-test")
    if r["adn_susana"] is None: f.append("diagnóstico ADN Susana")
    return f


def excel(rows, entidad) -> bytes:
    wb = Workbook()
    rs = wb.active; rs.title = "Resumen"
    det = wb.create_sheet("Detalle")
    cab = ["Cédula", "Nombre", "Cargo", "Proceso", "Fecha de asistencia", "Hora de asistencia",
           "Post-test", "ADN Susana", "Estado", "Qué le falta"]
    det.append(cab)
    cert = no = 0
    por_motivo = {}
    cumple = {"asistencia": 0, "post-test": 0, "diagnóstico ADN Susana": 0}
    for r in rows:
        fl = faltas(r)
        for k in cumple:
            if k not in fl: cumple[k] += 1
        ok = not fl
        cert += ok; no += not ok
        if fl: por_motivo[" + ".join(fl)] = por_motivo.get(" + ".join(fl), 0) + 1
        d, h = fasis(r["asistencia"]) if r["asistencia"] else ("Sin asistencia", "Sin asistencia")
        det.append([r["cedula"], r["nombre"], r["cargo"] or "", r["proceso"] or "", d, h,
                    f"{round(r['puntaje_posttest'] * 100 / r['total_posttest'])}%" if r["puntaje_posttest"] is not None else "Sin post-test",
                    f"{r['adn_susana']}%" if r["adn_susana"] is not None else "Sin diagnóstico",
                    "Certificado" if ok else "No certificado", "—" if ok else "Falta: " + ", ".join(fl)])
    hf = PatternFill("solid", fgColor="327531")
    for c in det[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hf
        c.alignment = Alignment(horizontal="center", vertical="center")
    verde, rojo = PatternFill("solid", fgColor="E3F2D3"), PatternFill("solid", fgColor="FBE3E3")
    for row in det.iter_rows(min_row=2):
        row[0].number_format = "0"
        fill = verde if row[8].value == "Certificado" else rojo
        row[8].fill = fill; row[9].fill = fill
    for i, h in enumerate(cab, 1):
        w = max(len(str(c.value or "")) for c in det[get_column_letter(i)]) + 3
        det.column_dimensions[get_column_letter(i)].width = min(max(w * 1.15, len(h) * 1.3 + 4), 50)
    det.auto_filter.ref = det.dimensions
    det.freeze_panes = "A2"

    def titulo(cell, texto):
        cell.value = texto; cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hf
    rs["A1"] = f"Certificación Reinducción 2026 — {entidad}"; rs["A1"].font = Font(bold=True, size=14, color="327531")
    rs["A2"] = "Se certifica a quien tiene asistencia, diagnóstico ADN Susana y post-test."
    titulo(rs["A4"], "Resultado"); titulo(rs["B4"], "Personas")
    for i, (k, v) in enumerate([("Total de la entidad", len(rows)), ("Completaron (certificados)", cert), ("No completaron", no)], 5):
        rs.cell(i, 1, k); rs.cell(i, 2, v)
    rs["A5"].font = rs["A6"].font = rs["A7"].font = Font(bold=True)
    titulo(rs["A9"], "Requisito"); titulo(rs["B9"], "Lo cumplen"); titulo(rs["C9"], "No lo cumplen")
    for i, (k, v) in enumerate(cumple.items(), 10):
        rs.cell(i, 1, k[0].upper() + k[1:]); rs.cell(i, 2, v); rs.cell(i, 3, len(rows) - v)
    titulo(rs["A14"], "Por qué no se certificó (lo que le falta)"); titulo(rs["B14"], "Personas")
    for i, (k, v) in enumerate(sorted(por_motivo.items(), key=lambda x: -x[1]) or [("Nadie: todos completaron", 0)], 15):
        rs.cell(i, 1, k[0].upper() + k[1:]); rs.cell(i, 2, v)
    rs.column_dimensions["A"].width = 52; rs.column_dimensions["B"].width = 14; rs.column_dimensions["C"].width = 16
    for row in rs.iter_rows(min_row=4):
        for c in row[1:3]: c.alignment = Alignment(horizontal="center")
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


async def main():
    entidad, salida = sys.argv[1], Path(sys.argv[2])
    salida.mkdir(parents=True, exist_ok=True)
    os.chdir(BACKEND)
    c = await asyncpg.connect(dotenv_values(".env")["DATABASE_URL"], statement_cache_size=0)
    rows = await c.fetch("""select cedula, nombre, cargo, proceso, asistencia, puntaje_posttest, total_posttest, adn_susana
                            from certificacion where upper(coalesce(entidad,'(SIN ENTIDAD)')) = any($1::text[]) order by nombre""",
                         [e.upper() for e in entidad.split("|")])
    etiqueta = os.environ.get("ETIQUETA") or re.sub(r"\b(Sas|Ltda|Sa)\b", lambda m: m.group(1).upper(), entidad.split("|")[0].title())
    base = marca.fondo_sin_marca(PLANTILLA)
    interno = io.BytesIO()
    n = 0
    with zipfile.ZipFile(interno, "w", zipfile.ZIP_STORED) as zc:   # los PDF ya vienen comprimidos
        for r in rows:
            if faltas(r): continue
            idn = await c.fetchval("""insert into certificados_emitidos (cedula, nombre, entidad) values ($1, $2, $3)
                                      on conflict (cedula) do update set nombre = excluded.nombre, entidad = excluded.entidad
                                      returning id""", r["cedula"], r["nombre"], etiqueta)
            zc.writestr(f"{r['cedula']}_{ascii_nombre(r['nombre'])}_reinduccion2026.pdf",
                        pdf_certificado(marca.marcar_jpeg(base, idn), r["nombre"], r["cedula"]))
            n += 1
    await c.close()
    xlsx = excel(rows, etiqueta)
    destino = salida / f"{etiqueta} - Certificacion Reinduccion2026.zip"
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_STORED) as z:
        z.writestr(f"Resumen {etiqueta}.xlsx", xlsx)
        z.writestr("certificados.zip", interno.getvalue())
    print(f"{etiqueta}: {len(rows)} personas, {n} certificados -> {destino} ({destino.stat().st_size/1e6:.1f} MB)")

asyncio.run(main())
